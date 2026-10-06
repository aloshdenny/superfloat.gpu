#!/usr/bin/env python3
"""
Rewrite section 6 (and the summary's v3 bullet) of the shuttle-plan Word
document with the measured 2 nm numbers (docs/v3_gt2n.md) and the v3_perf.py
workload results. The original is copied to "<name> (before 2 nm measured).docx"
first. Tables copy the document's own table format.

  python3 update_doc.py <doc.docx> <results/v3_perf.json>
"""
import copy
import json
import os
import shutil
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC, RES = sys.argv[1], sys.argv[2]
r = json.load(open(RES))
F = ["sf16", "fp16", "bf16"]
P, W = r["planned"], r["wide_vpu"]
W4 = r["wide_vpu_wload4"]


def fmt(v, nd=0):
    return f"{v:,.{nd}f}"


def x(a, b):
    return f"{a / b:.2f}×"


# ------------------------------------------------------------ doc helpers
d = docx.Document(DOC)
body = d.element.body
paras = d.paragraphs
start = next(p for p in paras if p.style.name == "Heading 1" and p.text.startswith("6. "))
end = next(p for p in paras if p.style.name == "Heading 1" and p.text.startswith("7. "))
el = start._p.getnext()
while el is not None and el is not end._p:
    nxt = el.getnext()
    body.remove(el)
    el = nxt
start.text = "6. The 2 nm GT2N chip: measured compute, on-die memory and large workloads"
anchor = end._p


def para(text, style="Normal", lead=None):
    p = d.add_paragraph(style=style)
    if lead:
        p.add_run(lead).bold = True
        p.add_run(" " + text)
    else:
        p.add_run(text)
    anchor.addprevious(p._p)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), fill)
    tcpr.append(s)


def table(rows, widths=None):
    n = len(rows[0])
    widths = widths or [2948] + [1077] * (n - 1)
    t = d.add_table(rows=len(rows), cols=n)
    t.style = d.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            c = t.cell(i, j)
            c.width = docx.shared.Twips(widths[j])
            c.paragraphs[0].text = ""
            run = c.paragraphs[0].add_run(str(text))
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, "1F3A5F")
            elif i % 2 == 0:
                shade(c, "EEF2F7")
    anchor.addprevious(t._tbl)
    para("")
    return t


def heading(text):
    para(text, "Heading 2")


# ---------------------------------------------------------------- section 6
para("This section replaces the earlier projection with measurements. The v3 processing elements were laid out "
     "on the GT2N 2 nm predictive PDK (OpenROAD flow, typical corner, 0.7 V) and powered with real GPT-2 data. A 2 nm "
     "die is then filled with as much on-die SRAM as fits, the same in every format, and run on GPT-2 and YOLOv8 "
     "workloads far larger than section 3's, with a performance model on the research server.")

heading("6.1 What was measured")
para("8 × 8 blocks of v3 processing elements (weight-stationary, as in the 128 × 128 matrix units) in each format, "
     "taken through synthesis, placement, clock tree and global routing on GT2N. Each is bit-exact against a Python "
     "model, and its gate-level netlist matches its RTL.", "List Bullet", "Blocks.")
para("The cheapest that keeps GPT-2's perplexity within 1% of FP32: a 24-bit fixed-point accumulator for SF16 (+0.55%), "
     "a 24-bit float for FP16 (+0.02%) and BF16 (+0.72%). The v2 PE's 16-bit SF16 accumulator fails (+609%).",
     "List Bullet", "Accumulators.")
para("A truncated multiplier (SF16 and FP16; BF16's product has nothing to drop) and clock gating of the weight "
     "registers (all three).", "List Bullet", "Every lever in every format it applies to.")
para("Gate-level switching activity on GPT-2 layer 5's real operands, at 2 GHz.", "List Bullet", "Power.")
para("SRAM (GT2N has no SRAM compiler), the vector units, wires beyond the block, and extracted parasitics (GT2N has "
     "no extraction rules; timing and power use the global route's estimate).", "List Bullet", "Not measured.")

heading("6.2 Raw on-die compute (measured)")
c = {f: r["planned"]["chip"][f] for f in F}
fm = r["formats"]
rows = [["Metric", "SF16", "FP16", "BF16", "SF16 vs FP16", "SF16 vs BF16"],
        ["Accumulator", "24-bit fixed", "24-bit float", "24-bit float", "", ""],
        ["Pipeline stages on the partial-sum path", "1", "3", "3", "", ""],
        ["Area per MAC (µm²)"] + [fmt(fm[f]["area_um2"], 1) for f in F] +
        [x(fm["fp16"]["area_um2"], fm["sf16"]["area_um2"]) + " less", x(fm["bf16"]["area_um2"], fm["sf16"]["area_um2"]) + " less"],
        ["Energy per MAC at 2 GHz (fJ)"] + [fmt(fm[f]["e_fj"]) for f in F] +
        [x(fm["fp16"]["e_fj"], fm["sf16"]["e_fj"]) + " less", x(fm["bf16"]["e_fj"], fm["sf16"]["e_fj"]) + " less"],
        ["Compute per mm² at 2 GHz (16-bit TFLOPS)"] + [fmt(c[f]["tflops_per_mm2_array"], 1) for f in F] +
        [x(c["sf16"]["tflops_per_mm2_array"], c["fp16"]["tflops_per_mm2_array"]), x(c["sf16"]["tflops_per_mm2_array"], c["bf16"]["tflops_per_mm2_array"])],
        ["Compute per watt (16-bit TFLOPS/W)"] + [fmt(c[f]["tflops_per_w"], 2) for f in F] +
        [x(c["sf16"]["tflops_per_w"], c["fp16"]["tflops_per_w"]), x(c["sf16"]["tflops_per_w"], c["bf16"]["tflops_per_w"])],
        ["Maximum clock (GHz)"] + [fmt(fm[f]["fmax"], 2) for f in F] +
        [x(fm["sf16"]["fmax"], fm["fp16"]["fmax"]), x(fm["sf16"]["fmax"], fm["bf16"]["fmax"])],
        ["GPT-2 perplexity vs FP32"] + [f"+{fm[f]['ppl_pct']:.2f}%" for f in F] + ["", ""]]
table(rows)
para("A MAC is two 16-bit FLOPs. All three run at 2 GHz; FP16 and BF16 need two more pipeline stages to get there, "
     "and with them reach a slightly higher maximum clock. SF16 delivers 2.53× FP16's and 2.09× BF16's compute per "
     "watt, and 1.70× and 1.48× their compute per mm².")

heading("6.3 The 2 nm die, with as much on-die memory as fits (modeled)")
cs = r["planned"]["chip"]["sf16"]
table([["Item", "Every format"],
       ["Die", "800 mm², near the 858 mm² reticle limit"],
       ["Matrix units (128 × 128)", "70 mm², the same area in every format"],
       ["Vector units, control cores, DMA, network", "14 mm²"],
       ["I/O and HBM PHYs", "40 mm²"],
       ["On-die SRAM", f"676 mm² at 38.1 Mb/mm²: {cs['sram_mb'] / 1e3:.2f} GB"],
       ["Off-die memory, for what does not fit", "4 HBM3E stacks: 4.8 TB/s, 144 GB"],
       ["Clock, power cap", "2 GHz, 700 W"]], widths=[3500, 5100])
table([["Per format", "SF16", "FP16", "BF16", "SF16 vs FP16", "SF16 vs BF16"],
       ["Matrix units"] + [str(c[f]["mxus"]) for f in F] + [x(c["sf16"]["mxus"], c["fp16"]["mxus"]), x(c["sf16"]["mxus"], c["bf16"]["mxus"])],
       ["MACs (millions)"] + [fmt(c[f]["macs"] / 1e6, 2) for f in F] + ["", ""],
       ["Peak 16-bit TFLOPS"] + [fmt(c[f]["peak_tflops"]) for f in F] +
       [x(c["sf16"]["peak_tflops"], c["fp16"]["peak_tflops"]), x(c["sf16"]["peak_tflops"], c["bf16"]["peak_tflops"])],
       ["Matrix-unit power at peak (W)"] + [fmt(c[f]["peak_w"]) for f in F] + ["", ""]])
para("The SRAM density is TSMC N2's published macro figure (section 5.3), and the die split and HBM are modeled; "
     "only the matrix units come from section 6.2. The model (v3/perf/v3_perf.py) runs on the research server. "
     "It is a roofline over four limits, and at batch 1 the work of each step is sequential:")
para("every GEMM is cut into 128 × 128 weight tiles. A tile takes the larger of its row count and 128 cycles, "
     "since its 128 weight rows load one per cycle while the previous tile computes. Each GEMM also pays its fill "
     "latency, 128 × stages + 128 cycles.", "List Bullet", "Matrix units:")
para("exp, GELU, softmax, SiLU and LayerNorm are lowered in software, at 6–20 ops per element. SF16 adds 2 ops per "
     "matmul input for its per-tensor scale. The plan's vector units are 16 × 128 lanes; the model also tries 65,536 "
     "lanes (16 × 4,096), whose area is not measured.", "List Bullet", "Vector units:")
para("whatever does not fit the SRAM, weights first, comes from HBM on every pass.", "List Bullet", "Memory:")
para("arithmetic energy at the 700 W cap.", "List Bullet", "Power:")


def g(sc, model, phase, B, T=None):
    for w in sc["gpt2"]:
        if w["model"] == model and w["phase"] == phase and w["batch"] == B and (T is None or w.get("tokens") == T):
            return w
    raise KeyError((model, phase, B, T))


def y(sc, model, img, B):
    return next(w for w in sc["yolo"] if w["model"] == model and w["img"] == img and w["batch"] == B)


LIM = {"mxu": "matrix units", "vector": "vector units", "hbm": "HBM", "power": "power", "latency": "latency"}


def wrow(label, w, key, scale=1.0, nd=0, higher=True):
    xs = [w[f] for f in F]
    if any(v is None for v in xs):
        return [label, "does not fit", "does not fit", "does not fit", "", "", "capacity"]
    v = [xx[key] * scale for xx in xs]
    ratio = (lambda a, b: a / b) if higher else (lambda a, b: b / a)
    lims = sorted({LIM[xx["bottleneck"]] for xx in xs}, key=list(LIM.values()).index)
    return [label] + [fmt(vv, nd) for vv in v] + [f"{ratio(v[0], v[1]):.2f}×", f"{ratio(v[0], v[2]):.2f}×", " / ".join(lims)]


def gpt2_table(sc):
    rows = [["GPT-2 workload", "SF16", "FP16", "BF16", "vs FP16", "vs BF16", "Limit"]]
    rows.append(wrow("Small, 1,024-token prompt × 16 (ms)", g(sc, "small", "prefill", 16, 1024), "ttft_ms", 1, 2, False))
    rows.append(wrow("XL, 1,024-token prompt × 1 (ms)", g(sc, "xl", "prefill", 1, 1024), "ttft_ms", 1, 2, False))
    rows.append(wrow("XL, 1,024-token prompt × 256 (s)", g(sc, "xl", "prefill", 256, 1024), "ttft_ms", 1e-3, 3, False))
    rows.append(wrow("Small, decode, 1 sequence (tokens/s)", g(sc, "small", "decode", 1), "per_s"))
    rows.append(wrow("Small, decode, 128 sequences (tokens/s)", g(sc, "small", "decode", 128), "per_s"))
    rows.append(wrow("Small, decode, 1,024 sequences (tokens/s)", g(sc, "small", "decode", 1024), "per_s"))
    rows.append(wrow("XL, decode, 1 sequence (tokens/s)", g(sc, "xl", "decode", 1), "per_s"))
    rows.append(wrow("XL, decode, 16 sequences (tokens/s)", g(sc, "xl", "decode", 16), "per_s"))
    rows.append(wrow("XL, decode, 1,024 sequences", g(sc, "xl", "decode", 1024), "per_s"))
    table(rows, widths=[2800, 950, 950, 950, 800, 800, 1350])


def yolo_table(sc):
    rows = [["YOLOv8 workload (images/s)", "SF16", "FP16", "BF16", "vs FP16", "vs BF16", "Limit"]]
    for m, img, B in [("yolov8n", 640, 1), ("yolov8n", 640, 1024), ("yolov8s", 640, 1), ("yolov8s", 640, 1024),
                      ("yolov8x", 1280, 1), ("yolov8x", 1280, 1024)]:
        rows.append(wrow(f"YOLOv8{m[-1]}, {img} × {img}, batch {B:,}", y(sc, m, img, B), "per_s"))
    table(rows, widths=[2800, 950, 950, 950, 800, 800, 1350])


heading("6.4 GPT-2, small to XL")
para("Prompts and decode at a 1,024-token context; ratios above 1× favour SF16. With the planned vector units:")
gpt2_table(P)
para("With 65,536 vector lanes:")
gpt2_table(W)

heading("6.5 YOLOv8, n to x")
para("With the planned vector units:")
yolo_table(P)
para("With 65,536 vector lanes:")
yolo_table(W)

heading("6.6 Where the bottleneck is")
def rng(vals, nd=2, suf="×"):
    return f"{min(vals):.{nd}f}–{max(vals):.{nd}f}{suf}"


big = [w for w in P["gpt2"] if w["phase"] == "prefill" and w["batch"] >= 16] + \
      [w for w in P["yolo"] if w["batch"] == 1024]
bigw = [w for w in W["gpt2"] if w["phase"] == "prefill" and w["batch"] >= 16]
yw_all = [w for w in W["yolo"] if w["batch"] == 1024]
util_p = [100 * w[f]["mxu_util"] for w in big for f in F]
trail = [100 * (1 - w["sf16"]["per_s"] / w["fp16"]["per_s"]) for w in big]
vs = lambda ws, o: [w["sf16"]["per_s"] / w[o]["per_s"] for w in ws]          # noqa: E731
xlw = g(W, "xl", "prefill", 256, 1024)
para(f"With the planned 2,048 lanes, every batch of prompts and images waits on the lowered exp, GELU, softmax and "
     f"SiLU. The matrix units are busy only {min(util_p):.0f}–{max(util_p):.0f}% of the time, and SF16 trails by "
     f"{min(trail):.0f}–{max(trail):.0f}%, because its per-tensor scaling adds vector ops. At 65,536 lanes the matrix "
     f"units become the limit ({100 * xlw['sf16']['mxu_util']:.0f}% busy on GPT-2 XL prompts). SF16 then leads "
     f"FP16 by {rng(vs(bigw, 'fp16'))} on GPT-2 prompts and {rng(vs(yw_all, 'fp16'))} on YOLOv8 batches, and BF16 by "
     f"{rng(vs(bigw, 'bf16'))} and {rng(vs(yw_all, 'bf16'))}: around its 1.73× and 1.49× more matrix units. The spread is "
     f"how evenly each model's weight tiles divide among 64, 37 or 43 units.",
     "List Bullet", "Vector units first.")
s1 = g(W, "small", "decode", 1)
s1w = g(W4, "small", "decode", 1)
para(f"One sequence is latency-bound: each 128 × 128 weight tile takes 128 cycles to load and serves one token. SF16 "
     f"leads by {s1['sf16']['per_s'] / s1['fp16']['per_s']:.2f}× and {s1['sf16']['per_s'] / s1['bf16']['per_s']:.2f}× (more matrix units, one pipeline "
     f"stage instead of three): {1e6 / s1['sf16']['per_s']:.1f} µs per GPT-2 small token at a 1,024-token context. Loading 4 weight rows "
     f"per cycle gives {1e6 / s1w['sf16']['per_s']:.1f} µs and {s1w['sf16']['per_s'] / s1w['fp16']['per_s']:.2f}× / "
     f"{s1w['sf16']['per_s'] / s1w['bf16']['per_s']:.2f}×.", "List Bullet", "Decode next: the weight load.")
nl, dd, _ = (48, 1600, 25)
xl_w = 2 * (nl * 12 * dd * dd + 50257 * dd + 1024 * dd)
kv_seq = 2 * nl * dd * 2 * 1024
xl16, xl128 = g(W, "xl", "decode", 16), g(W, "xl", "decode", 128)
para(f"The {cs['sram_mb'] / 1e3:.2f} GB of SRAM holds GPT-2 XL's {xl_w / 1e9:.2f} GB of weights with "
     f"{(cs['sram_mb'] * 1e6 - xl_w) / 1e6:.0f} MB to spare. Its KV cache ({kv_seq / 1e6:.0f} MB per 1,024-token sequence) "
     f"goes to HBM. XL decode at 16 sequences or more is therefore HBM-bound, at {fmt(xl128['sf16']['per_s'])}–"
     f"{fmt(xl16['sf16']['per_s'])} tokens/s in every format. 1,024 XL sequences ({1024 * kv_seq / 1e9:.0f} GB of KV "
     f"cache) fit neither the die nor HBM. GPT-2 small ties at 1,024 sequences for the same reason.",
     "List Bullet", "Then memory capacity.")
umax16 = max(w["fp16"]["mxu_util"] for w in W["gpt2"] + W["yolo"] if w["fp16"])
para(f"Power is never the limit at 700 W in these runs. SF16's arrays peak at {fmt(c['sf16']['peak_w'])} W. FP16's would "
     f"need {fmt(c['fp16']['peak_w'])} W at full rate but are never more than {100 * umax16:.0f}% busy.", "List Bullet", "Power.")
para("SF16's measured advantage shows wherever the matrix units are the limit. To get there, v3 needs three changes "
     "that apply to all three formats: much wider vector units (or a few hardware special functions), faster weight "
     "loading for decode, and more capacity for long-context batch decode.")

heading("6.7 Against the push gate")
vw = W["v1"]
yw = y(W, "yolov8n", 640, 1024)
table([["Track", "Measure", "SF16", "FP16", "BF16", "vs FP16", "vs BF16"],
       ["Area", "TFLOPS per mm² (measured)"] + [fmt(c[f]["tflops_per_mm2_array"], 1) for f in F] +
       [x(c["sf16"]["tflops_per_mm2_array"], c["fp16"]["tflops_per_mm2_array"]), x(c["sf16"]["tflops_per_mm2_array"], c["bf16"]["tflops_per_mm2_array"])],
       ["Power", "Energy per MAC, fJ (measured)"] + [fmt(fm[f]["e_fj"]) for f in F] +
       [x(fm["fp16"]["e_fj"], fm["sf16"]["e_fj"]) + " less", x(fm["bf16"]["e_fj"], fm["sf16"]["e_fj"]) + " less"],
       ["FLOPS", "Peak TFLOPS in 70 mm² of arrays"] + [fmt(c[f]["peak_tflops"]) for f in F] +
       [x(c["sf16"]["peak_tflops"], c["fp16"]["peak_tflops"]), x(c["sf16"]["peak_tflops"], c["bf16"]["peak_tflops"])],
       ["Time to first token", "GPT-2 small, 32-token prompt (µs)"] + [fmt(vw[f]["ttft_us"], 1) for f in F] +
       [x(vw["fp16"]["ttft_us"], vw["sf16"]["ttft_us"]) + " sooner", x(vw["bf16"]["ttft_us"], vw["sf16"]["ttft_us"]) + " sooner"],
       ["Generation", "GPT-2 small, µs per token"] + [fmt(vw[f]["tpot_us"], 1) for f in F] +
       [x(vw["fp16"]["tpot_us"], vw["sf16"]["tpot_us"]) + " faster", x(vw["bf16"]["tpot_us"], vw["sf16"]["tpot_us"]) + " faster"],
       ["CNN", "YOLOv8n images/s, batch 1,024"] + [fmt(yw[f]["per_s"]) for f in F] +
       [x(yw["sf16"]["per_s"], yw["fp16"]["per_s"]), x(yw["sf16"]["per_s"], yw["bf16"]["per_s"])]],
      widths=[1400, 2500, 900, 900, 900, 1000, 1000])
para("Workload rows use 65,536 vector lanes. SF16 leads on every track, but clears 2× against both formats only on "
     "power.")

heading("6.8 Against the v1 chip")
v1 = {"ttft": 140.2, "tpot": 21.3, "e2e": 800.0, "yolo": 205.0}
sf = vw["sf16"]
table([["SF16", "v1, measured (SKY130, pin bus)", "v3 at 2 nm (modeled, 65,536 lanes)", "Factor"],
       ["Time to first token, 32-token prompt", "140.2 s", f"{sf['ttft_us']:.1f} µs", f"{v1['ttft'] / (sf['ttft_us'] * 1e-6) / 1e6:.1f} million×"],
       ["Time per output token", "21.3 s", f"{sf['tpot_us']:.1f} µs", f"{v1['tpot'] / (sf['tpot_us'] * 1e-6) / 1e6:.2f} million×"],
       ["End to end, 32 + 32 tokens", "800 s", f"{sf['e2e_us']:.0f} µs", f"{v1['e2e'] / (sf['e2e_us'] * 1e-6) / 1e6:.1f} million×"],
       ["YOLOv8n, time per image", "205 s", f"{sf['yolo_us']:.1f} µs", f"{v1['yolo'] / (sf['yolo_us'] * 1e-6) / 1e6:.1f} million×"]],
      widths=[2900, 2000, 2400, 1300])
para("The earlier projection's 2.0 µs per token assumed power-bound arrays. The model now includes each tile's "
     "128-cycle weight load, which sets batch-1 decode.")

# ----------------------------------------------------------- summary bullet
for p in d.paragraphs:
    if p.text.startswith("v3 at 2 nm"):
        for rr in p.runs[1:]:
            rr.text = ""
        p.runs[0].text = ("v3 at 2 nm, measured: on the GT2N 2 nm PDK, SF16 delivers 2.53× FP16's and 2.09× BF16's compute "
                          "per watt (8.6 TFLOPS/W) and 1.70× and 1.48× their compute per mm² (60 TFLOPS/mm²). On an 800 mm² "
                          "die with 3.2 GB of SRAM in every format, it leads GPT-2 and YOLOv8 by 1.4–1.7× wherever the matrix "
                          "units are the limit. The plan's vector units, decode's weight loading and KV-cache capacity are "
                          "the bottlenecks to fix (section 6).")
        break

base, ext = os.path.splitext(DOC)
backup = f"{base} (before 2 nm measured){ext}"
if not os.path.exists(backup):
    shutil.copy2(DOC, backup)
d.save(DOC)
print("saved", DOC, "| backup", backup)
